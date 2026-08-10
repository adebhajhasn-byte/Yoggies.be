"""
audit.py — verify the data and the built site
=============================================
Run this after every build. It is the check that the old pipeline never had:
187 untranslated descriptions shipped silently because nothing was looking.

Two halves:
  DATA  the master spreadsheet - translation, completeness, duplicates
  SITE  the generated HTML - SEO tags, images, links, purchase paths

Every check prints PASS, WARN or FAIL. The exit code is non-zero if anything
FAILs, so it can gate a build.

Usage
    python pipe/audit.py
    python pipe/audit.py --data-only
    python pipe/audit.py --verbose      # list offending items
"""

from __future__ import annotations

import argparse
import glob
import html
import os
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency. Run: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "READY_FOR_BUILD_FINAL.xlsx"
SHEET = "READY_FOR_BUILD"
SITE = ROOT / "site"

CZECH = re.compile(
    r"\[(CZ|NL|EN|DE|FR|SK|PL)\]"
    r"|[ěščřůĚŠČŘŮ]"
    r"|\b(jsou|které|pro psy|pro kočky|granule|krmivo|balení|masem|lisované|Kč)\b",
    re.IGNORECASE,
)
# Brand-ish exceptions that legitimately keep diacritics, plus the handful of
# Czech PERSONAL NAMES that appear as credits on the site. Named individually on
# purpose: whitelisting the diacritics themselves would blind the check to real
# untranslated Czech.
ALLOWED = re.compile(
    r"Ibérico|Yoggies|B\.A\.R\.F\.|Markus-Mühle"
    r"|Jan\s+Herčík|Markéta\s+Reinert|Daniela\s+Králová",
    re.IGNORECASE,
)

# Code that leaked into prose. 111 descriptions once opened with
# ".yoggies-section { max-width: 1100px; ... }" because html_to_text() stripped
# <style> tags but kept the stylesheet between them. Each alternative is a shape
# prose does not have: a declaration block, an at-rule, a JS function or a var.
CSS_JS = re.compile(
    r"\{[^{}]*[A-Za-z-]+\s*:\s*[^{};]+[;}]"
    r"|@(?:media|import|font-face|keyframes|supports)\b"
    r"|\bfunction\s*\("
    r"|\bvar\s+\w+\s*=",
    re.IGNORECASE,
)

results: list[tuple[str, str, str]] = []
VERBOSE = False


def check(name: str, ok: bool, detail: str = "", warn_only: bool = False) -> None:
    level = "PASS" if ok else ("WARN" if warn_only else "FAIL")
    results.append((level, name, detail))
    symbol = {"PASS": "  ok  ", "WARN": " warn ", "FAIL": " FAIL "}[level]
    print(f"[{symbol}] {name}" + (f"  -  {detail}" if detail else ""))


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:70]


def audit_data() -> None:
    print("\n=== DATA ===")
    if not MASTER.exists():
        check("master file exists", False, str(MASTER.name))
        return

    wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = {str(h): i for i, h in enumerate(next(rows)) if h}
    data = [r for r in rows if any(r)]
    wb.close()

    def col(row, name):
        i = header.get(name)
        return row[i] if i is not None and i < len(row) else None

    total = len(data)
    check("products present", total > 0, f"{total} rows")

    untranslated = [r for r in data
                    if col(r, "description_nl")
                    and CZECH.search(ALLOWED.sub("", str(col(r, "description_nl"))))]
    check("descriptions fully translated", not untranslated,
          f"{len(untranslated)} of {total} still contain Czech")
    if VERBOSE and untranslated:
        for r in untranslated[:10]:
            print(f"          - {str(col(r,'name_nl'))[:70]}")

    cz_names = [r for r in data if CZECH.search(ALLOWED.sub("", str(col(r, "name_nl") or "")))]
    check("names fully translated", not cz_names, f"{len(cz_names)} names contain Czech")

    empty_desc = [r for r in data if not str(col(r, "description_nl") or "").strip()]
    check("every product has a description", not empty_desc, f"{len(empty_desc)} empty")

    code_desc = [r for r in data
                 for field in ("description_nl", "description_cz")
                 if CSS_JS.search(str(col(r, field) or ""))]
    code_desc = list({id(r): r for r in code_desc}.values())
    check("no CSS or JavaScript in descriptions", not code_desc,
          f"{len(code_desc)} of {total} contain stylesheet or script text")
    if VERBOSE and code_desc:
        for r in code_desc[:10]:
            snippet = " ".join(str(col(r, "description_nl") or "").split())[:60]
            print(f"          - {str(col(r,'name_nl'))[:46]}  |  {snippet}")

    no_img = [r for r in data if not col(r, "image_main")]
    check("every product has a main image", not no_img, f"{len(no_img)} missing")

    names = Counter(str(col(r, "name_nl") or "") for r in data)
    dupes = [k for k, v in names.items() if v > 1 and k]
    check("no duplicate product names", not dupes, f"{len(dupes)} duplicated")

    slugs = Counter(slugify(col(r, "name_nl")) for r in data)
    collisions = [k for k, v in slugs.items() if v > 1]
    check("no URL collisions", not collisions, f"{len(collisions)} slugs used twice")

    bad_calc = [r for r in data if col(r, "has_calculator") and not col(r, "dosage_cz")]
    check("calculator only where dosage data exists", not bad_calc,
          f"{len(bad_calc)} have the calculator on without dosage", warn_only=True)

    kc = [r for r in data
          if "kč" in f"{col(r,'name_nl')} {col(r,'description_nl')}".lower()]
    check("no Czech currency in Dutch text", not kc, f"{len(kc)} mention Kc")

    statuses = Counter(str(col(r, "match_status") or "?") for r in data)
    linked = sum(v for k, v in statuses.items() if k.startswith("MATCH"))
    check("products linked to DogLife", linked > 0,
          f"{linked} linked, {statuses.get('NOT_ON_DOGLIFE', 0)} information-only",
          warn_only=True)

    review = statuses.get("REVIEW", 0)
    check("no pending match reviews", review == 0,
          f"{review} awaiting your decision in data/review_matches.xlsx",
          warn_only=True)


def audit_site() -> None:
    print("\n=== SITE ===")
    if not SITE.exists():
        check("site folder exists", False, "run yoggies_pipeline.py first")
        return

    pages = glob.glob(str(SITE / "**" / "*.html"), recursive=True)
    check("pages generated", bool(pages), f"{len(pages)} html files")
    if not pages:
        return

    no_desc = no_canon = onclick = czech_visible = code_visible = 0
    noalt = imgs = 0
    hotlinks: Counter = Counter()
    broken: Counter = Counter()
    code_pages: list[str] = []

    for path in pages:
        text = Path(path).read_text(encoding="utf-8", errors="replace")
        folder = os.path.dirname(path)

        if 'name="description"' not in text:
            no_desc += 1
        if 'rel="canonical"' not in text:
            no_canon += 1
        if 'onclick="window.location' in text:
            onclick += 1

        visible = re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ",
                         text, flags=re.S)
        if CZECH.search(ALLOWED.sub("", visible)):
            czech_visible += 1
        # Real <style>/<script> blocks are stripped above, so a hit here is a
        # stylesheet that arrived as description text and is rendering as prose.
        if CSS_JS.search(html.unescape(visible)):
            code_visible += 1
            code_pages.append(os.path.basename(path))

        for tag in re.findall(r"<img[^>]*>", text):
            imgs += 1
            if "alt=" not in tag:
                noalt += 1

        # Hotlinking is about IMAGE sources only. Outbound <a href> links to
        # doglife.be are the purchase links - those are supposed to be there.
        for src in re.findall(r'<img[^>]+src="([^"]+)"', text):
            src = html.unescape(src)
            if src.startswith("http"):
                hotlinks[re.sub(r"https?://([^/]+).*", r"\1", src)] += 1

        for url in re.findall(r'(?:href|src)="([^"]+)"', text):
            url = html.unescape(url)
            if url.startswith(("http", "#", "mailto:", "tel:", "data:", "javascript:")):
                continue
            target = os.path.normpath(
                os.path.join(folder, url.split("#")[0].split("?")[0]))
            if target and "'" not in target and not os.path.exists(target):
                broken[target] += 1

    check("every page has a meta description", no_desc == 0, f"{no_desc} missing")
    check("every page has a canonical URL", no_canon == 0, f"{no_canon} missing")
    check("all product cards are real links", onclick == 0, f"{onclick} still use onclick")
    check("every image has alt text", noalt == 0, f"{noalt} of {imgs} missing")
    check("no images loaded from other sites", not hotlinks, f"{dict(hotlinks)}")
    check("no broken internal links", not broken, f"{len(broken)} broken paths")
    if VERBOSE and broken:
        for k, v in broken.most_common(10):
            print(f"          - {v}x  {k}")
    check("no Czech text visible on pages", czech_visible == 0,
          f"{czech_visible} pages show Czech")
    check("no CSS or JavaScript visible on pages", code_visible == 0,
          f"{code_visible} pages render code as text")
    if VERBOSE and code_pages:
        for name in code_pages[:10]:
            print(f"          - {name}")

    for required in ("sitemap.xml", "robots.txt"):
        check(f"{required} generated", (SITE / required).exists())

    products = glob.glob(str(SITE / "products" / "*.html"))
    if products:
        # A product row that shares a slug with another row loses its page
        # silently: the file is simply written twice. Compare what the builder
        # says it published against the real (non-redirect) pages on disk.
        stubs = [p for p in products
                 if 'http-equiv="refresh"' in Path(p).read_text(
                     encoding="utf-8", errors="replace")]
        real = len(products) - len(stubs)
        manifest_path = ROOT / "data" / "build_manifest.json"
        if manifest_path.exists():
            import json
            m = json.loads(manifest_path.read_text(encoding="utf-8"))
            rows = m.get("product_rows", 0)
            check("every product row has its own page", rows == real,
                  f"{rows} product rows, {real} product pages"
                  + ("" if rows == real else f" - {rows - real} overwritten"))
        else:
            check("every product row has its own page", False,
                  "data/build_manifest.json missing - rebuild the site",
                  warn_only=True)

        both = neither = 0
        for path in products:
            text = Path(path).read_text(encoding="utf-8", errors="replace")
            has_msg = "niet beschikbaar" in text.lower()
            has_link = bool(re.search(r"doglife\.be/product/", text))
            if has_msg and has_link:
                both += 1
            elif not has_msg and not has_link:
                neither += 1
        check("no contradictory purchase blocks", both == 0,
              f"{both} pages show a buy button AND 'not available'")
        check("every product page has a call to action", neither == 0,
              f"{neither} pages show neither", warn_only=True)


def main() -> int:
    global VERBOSE
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-only", action="store_true")
    ap.add_argument("--site-only", action="store_true")
    ap.add_argument("--verbose", action="store_true")
    args = ap.parse_args()
    VERBOSE = args.verbose

    if not args.site_only:
        audit_data()
    if not args.data_only:
        audit_site()

    failed = [r for r in results if r[0] == "FAIL"]
    warned = [r for r in results if r[0] == "WARN"]
    passed = [r for r in results if r[0] == "PASS"]

    print(f"\n{'=' * 52}")
    print(f"  {len(passed)} passed   {len(warned)} warnings   {len(failed)} failed")
    if failed:
        print("\nMust fix:")
        for _, name, detail in failed:
            print(f"  - {name}: {detail}")
    print("=" * 52)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
