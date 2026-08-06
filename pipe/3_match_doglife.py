"""
3_match_doglife.py — link products to their DogLife purchase page
=================================================================
The old matching accepted only character-identical names. Every one of its
122 matches scored exactly 1.00 and everything else was dumped into
NOT_ON_DOGLIFE, because "1,5 kg Yoggies BARF hovezi a kruti" is not literally
equal to "1,5 kg (6x250gr) Yoggies BARF Rund en Kalkoen Compleet".

This matches on what actually identifies a product: pack size and flavour,
with a Czech/Dutch vocabulary so krůta and kalkoen compare equal. Rare words
count for more than common ones, so "Chlorella" outweighs "voor honden".

Three outcomes, never a silent guess:
  MATCH            confident, used automatically
  REVIEW           plausible, written to review_matches.xlsx for you to judge
  NOT_ON_DOGLIFE   no candidate - the product is information-only

Output
  data/doglife_matches.json   handle -> match record
  data/review_matches.xlsx    put y or n in the `keep` column
  data/doglife_unused.xlsx    DogLife products that are on no page yet

Review loop
  1. python pipe/3_match_doglife.py
  2. open data/review_matches.xlsx, fill in `keep`
  3. python pipe/3_match_doglife.py --apply-review
Decisions are stored in data/match_decisions.json and always win over scoring.

Usage
    python pipe/3_match_doglife.py
    python pipe/3_match_doglife.py --apply-review
    python pipe/3_match_doglife.py --threshold 0.62
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency. Run: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
EXPORTS = DATA / "doglife"
CATALOGUE = DATA / "yoggies_translated.json"
FALLBACK = DATA / "yoggies_catalogue.json"
OUT = DATA / "doglife_matches.json"
REVIEW = DATA / "review_matches.xlsx"
UNUSED = DATA / "doglife_unused.xlsx"
DECISIONS = DATA / "match_decisions.json"

AUTO_ACCEPT = 0.70
REVIEW_FLOOR = 0.42

SYNONYMS = {
    "kalkoen": "turkey", "krůta": "turkey", "kruta": "turkey", "kruti": "turkey",
    "krůtí": "turkey", "krocan": "turkey", "krocani": "turkey", "kalkoenvlees": "turkey",
    "kip": "chicken", "kuře": "chicken", "kure": "chicken", "kureci": "chicken",
    "kuřecí": "chicken", "kippen": "chicken", "kippenvlees": "chicken",
    "rund": "beef", "rundvlees": "beef", "hovězí": "beef", "hovezi": "beef",
    "zalm": "salmon", "losos": "salmon", "lososovy": "salmon",
    "lam": "lamb", "jehněčí": "lamb", "jehneci": "lamb",
    "varken": "pork", "vepřové": "pork", "veprove": "pork", "varkensvlees": "pork",
    "eend": "duck", "kachna": "duck", "kachní": "duck", "kachni": "duck",
    "konijn": "rabbit", "králík": "rabbit", "kralik": "rabbit", "kralici": "rabbit",
    "geit": "goat", "koza": "goat", "kozí": "goat", "kozi": "goat",
    "hert": "venison", "wild": "venison", "zvěřina": "venison", "zverina": "venison",
    "jeleni": "venison", "witvis": "whitefish", "treska": "whitefish",
    "paard": "horse", "kůň": "horse", "konske": "horse",
    "kangoeroe": "kangaroo", "klokani": "kangaroo",
    "brokken": "kibble", "granule": "kibble", "korrels": "kibble",
    "minigranule": "minikibble", "minibrokken": "minikibble",
    "paté": "pate", "pate": "pate", "paštika": "pate", "pastika": "pate",
    "ragout": "ragu", "ragú": "ragu", "ragu": "ragu",
    "blik": "can", "konzerva": "can", "conserve": "can",
    "yoghurt": "yoghurt", "jogurt": "yoghurt",
    "snoepjes": "treat", "pamlsky": "treat", "pamlsek": "treat", "snack": "treat",
    "gedroogd": "dried", "sušené": "dried", "susene": "dried", "susena": "dried",
    "gekookt": "cooked", "vařené": "cooked", "varene": "cooked", "varena": "cooked",
    "olie": "oil", "olej": "oil", "oleje": "oil",
    "stelivo": "litter", "kattenbakvulling": "litter", "podestylka": "litter",
    "kattenvoer": "cat", "kočky": "cat", "kocky": "cat", "katten": "cat", "kat": "cat",
    "honden": "dog", "psy": "dog", "psa": "dog", "hond": "dog", "psu": "dog",
    "probiotica": "probiotic", "probiotiky": "probiotic",
    "monoproteine": "mono", "monoprotein": "mono", "monoproteïne": "mono",
    "gewrichten": "joints", "klouby": "joints", "kloubni": "joints",
}

STOP = {
    "yoggies", "voor", "pro", "met", "en", "a", "s", "de", "het", "een", "van",
    "compleet", "kompletni", "the", "of", "and", "in", "op", "je", "za", "se",
}


def strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(c))


def size_tokens(text: str) -> set[str]:
    """Normalise pack sizes so 1,5 kg == 1.5kg == 1500 g."""
    out: set[str] = set()
    t = text.lower().replace(",", ".")
    for count, each, unit in re.findall(
            r"(\d+)\s*[x×]\s*(\d+(?:\.\d+)?)\s*(kg|g|gr|ml|l)\b", t):
        grams = float(each) * (1000 if unit in ("kg", "l") else 1)
        out.add(f"p{int(count)}")
        out.add(f"u{int(grams)}")
        out.add(f"w{int(grams) * int(count)}")
    if not out:
        for value, unit in re.findall(r"(\d+(?:\.\d+)?)\s*(kg|g|gr|ml|l)\b", t):
            grams = float(value) * (1000 if unit in ("kg", "l") else 1)
            out.add(f"w{int(grams)}")
    return out


def word_tokens(text: str) -> set[str]:
    flat = strip_accents(text.lower())
    flat = re.sub(r"\d+(?:[.,]\d+)?\s*(kg|g|gr|ml|l)\b", " ", flat)
    flat = re.sub(r"[^a-z0-9]+", " ", flat)
    words = set()
    for w in flat.split():
        if len(w) < 2 or w in STOP or w.isdigit():
            continue
        words.add(SYNONYMS.get(w, w))
    return words


def load_exports() -> dict[str, dict]:
    catalogue: dict[str, dict] = {}
    files = [f for f in sorted(EXPORTS.glob("*.xlsx")) if not f.name.startswith("~$")]
    if not files:
        sys.exit(f"No DogLife exports in {EXPORTS}. Put your doglife*.xlsx files there.")

    for path in files:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = {str(h): i for i, h in enumerate(next(rows)) if h}

        def col(row, *names):
            for n in names:
                i = header.get(n)
                if i is not None and i < len(row) and row[i]:
                    return str(row[i]).strip()
            return ""

        added = 0
        for row in rows:
            title = col(row, "woocommerce-loop-product__title")
            url = col(row, "woocommerce-LoopProduct-link href", "button href").rstrip("/")
            if not title or "/product/" not in url:
                continue
            if url not in catalogue:
                added += 1
            catalogue[url] = {
                "title": title,
                "url": url,
                "image": col(row, "attachment-woocommerce_thumbnail src"),
                "price": col(row, "woocommerce-Price-amount", "price"),
                "excerpt": col(row, "woocommerce-product-excerpt"),
                "source": path.name,
            }
        wb.close()
        print(f"  {path.name}: +{added}")
    return catalogue


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--threshold", type=float, default=AUTO_ACCEPT)
    ap.add_argument("--apply-review", action="store_true")
    args = ap.parse_args()

    src = CATALOGUE if CATALOGUE.exists() else FALLBACK
    if not src.exists():
        sys.exit("No catalogue found. Run pipe/1_fetch_yoggies.py first.")
    products = json.loads(src.read_text(encoding="utf-8"))
    print(f"catalogue: {len(products)} products (from {src.name})\n")

    print("DogLife exports:")
    shop = load_exports()
    print(f"  {len(shop)} unique DogLife products\n")

    decisions = json.loads(DECISIONS.read_text(encoding="utf-8")) if DECISIONS.exists() else {}

    # Matches already verified in the existing master are worth keeping - they
    # were correct, just incomplete. Seed them so nothing regresses and you are
    # not asked to re-review pairs that were already right.
    master = ROOT / "READY_FOR_BUILD_FINAL.xlsx"
    if master.exists():
        wbm = openpyxl.load_workbook(master, read_only=True, data_only=True)
        wsm = wbm[wbm.sheetnames[0]]
        mrows = wsm.iter_rows(values_only=True)
        mh = {str(x): i for i, x in enumerate(next(mrows)) if x}
        seeded = 0
        if "yoggies_url" in mh and "doglife_url" in mh:
            for row in mrows:
                status = str(row[mh["match_status"]] or "") if "match_status" in mh else ""
                url = str(row[mh["doglife_url"]] or "").rstrip("/")
                handle = str(row[mh["yoggies_url"]] or "").rstrip("/").split("/")[-1]
                if handle and status.startswith("MATCH") and "/product/" in url:
                    if handle not in decisions and url in shop:
                        decisions[handle] = url
                        seeded += 1
        wbm.close()
        if seeded:
            print(f"carried over {seeded} verified matches from the existing master\n")

    if args.apply_review:
        if not REVIEW.exists():
            sys.exit(f"{REVIEW.name} not found - run without --apply-review first.")
        wb = openpyxl.load_workbook(REVIEW, data_only=True)
        ws = wb.active
        header = {str(c.value): i for i, c in enumerate(ws[1])}
        applied = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            keep = str(row[header["keep"]] or "").strip().lower()
            handle = str(row[header["handle"]] or "")
            if keep in ("y", "yes", "j", "ja", "1"):
                decisions[handle] = row[header["doglife_url"]]
                applied += 1
            elif keep in ("n", "no", "nee", "0"):
                decisions[handle] = None
                applied += 1
        DECISIONS.write_text(json.dumps(decisions, indent=1, ensure_ascii=False), encoding="utf-8")
        print(f"applied {applied} decisions\n")

    # Rare words are more informative than common ones.
    prepared = [(i, word_tokens(i["title"]), size_tokens(i["title"])) for i in shop.values()]
    df = Counter()
    for _, words, _ in prepared:
        df.update(words)
    total = max(1, len(prepared))

    def idf(token: str) -> float:
        return math.log((total + 1) / (df.get(token, 0) + 1)) + 1.0

    def score(aw, asz, bw, bsz) -> float:
        if not aw or not bw:
            return 0.0
        shared = sum(idf(t) for t in aw & bw)
        smaller = min(sum(idf(t) for t in aw), sum(idf(t) for t in bw)) or 1.0
        word_sim = shared / smaller
        size_sim = len(asz & bsz) / len(asz | bsz) if (asz and bsz) else 0.5
        result = 0.70 * word_sim + 0.30 * size_sim
        if asz and bsz and not (asz & bsz):
            result *= 0.45          # different pack size, almost certainly not it
        return result

    matches: dict[str, dict] = {}
    review, used = [], set()
    stats = Counter()

    for item in products:
        handle = item["handle"]
        name = item.get("name_nl") or item["name_cz"]

        if handle in decisions:
            url = decisions[handle]
            hit = shop.get(url) if url else None
            status = "MATCH_CONFIRMED" if hit else "NOT_ON_DOGLIFE"
            # A decided `n` is a decision too, but it is not a link - count it
            # with the information-only products or the summary overstates them.
            stats["confirmed" if hit else "none"] += 1
        else:
            aw, asz = word_tokens(name), size_tokens(name)
            best, best_score = None, 0.0
            for cand, bw, bsz in prepared:
                s = score(aw, asz, bw, bsz)
                if s > best_score:
                    best_score, best = s, cand
            if best and best_score >= args.threshold:
                hit, status = best, "MATCH"
                stats["match"] += 1
            elif best and best_score >= REVIEW_FLOOR:
                hit, status = None, "REVIEW"
                review.append([handle, name, best["title"], best["url"],
                               best["price"], round(best_score, 3), ""])
                stats["review"] += 1
            else:
                hit, status = None, "NOT_ON_DOGLIFE"
                stats["none"] += 1

        if hit:
            used.add(hit["url"])
        matches[handle] = {
            "status": status,
            "doglife_url": hit["url"] if hit else "",
            "doglife_image": hit["image"] if hit else "",
            "doglife_title": hit["title"] if hit else "",
            "doglife_price": hit["price"] if hit else "",
        }

    OUT.write_text(json.dumps(matches, ensure_ascii=False, indent=1), encoding="utf-8")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "REVIEW"
    ws.append(["handle", "our_name", "doglife_title", "doglife_url", "price", "score", "keep"])
    for r in sorted(review, key=lambda x: -x[5]):
        ws.append(r)
    ws.freeze_panes = "A2"
    wb.save(REVIEW)

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "NOT_ON_SITE"
    ws2.append(["doglife_title", "doglife_url", "price", "image", "excerpt", "source"])
    for item in shop.values():
        if item["url"] not in used:
            ws2.append([item["title"], item["url"], item["price"],
                        item["image"], item["excerpt"], item["source"]])
    ws2.freeze_panes = "A2"
    wb2.save(UNUSED)

    linked = stats["match"] + stats["confirmed"]
    print("Results")
    print(f"  linked to DogLife  : {linked}")
    print(f"  need your review   : {stats['review']}   -> {REVIEW.relative_to(ROOT)}")
    print(f"  information-only   : {stats['none']}")
    print(f"  DogLife not on site: {len(shop) - len(used)}   -> {UNUSED.relative_to(ROOT)}")
    print(f"\nWrote {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
