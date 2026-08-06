"""
4_build_master.py — assemble READY_FOR_BUILD_FINAL.xlsx
========================================================
Combines the three sources into the single file yoggies_pipeline.py consumes:

  data/yoggies_translated.json   names + descriptions (Czech source, Dutch text)
  data/doglife_matches.json      purchase links
  READY_FOR_BUILD_FINAL.xlsx     the EXISTING file - ingredients, dosages, icons,
                                 the hand-built calculator flags AND the Dutch
                                 name/description are read back out and preserved

Nothing hand-made is thrown away. Products are matched to the old file by
yoggies.cz URL first, then by Czech name, so work done over previous sessions
carries forward even where a product was renamed.

Translated text is expensive, so name_nl and description_nl resolve in this
order and a rebuild never destroys a translation just because the translation
step has not been re-run:

    1. yoggies_translated.json  (a fresh translation, when non-empty)
    2. the preserve source      (the translation from the previous master)
    3. the Czech source         (untranslated passthrough, last resort)

The old file is backed up before it is replaced.

Usage
    python pipe/4_build_master.py
    python pipe/4_build_master.py --dry-run     # report only, write nothing
    python pipe/4_build_master.py --preserve-from OLD.xlsx
                                                # preserve from a specific file
                                                # (e.g. a backup) instead of the
                                                # master about to be overwritten
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Missing dependency. Run: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
TRANSLATED = DATA / "yoggies_translated.json"
CATALOGUE = DATA / "yoggies_catalogue.json"
MATCHES = DATA / "doglife_matches.json"
MASTER = ROOT / "READY_FOR_BUILD_FINAL.xlsx"
SHEET = "READY_FOR_BUILD"

COLUMNS = [
    "row_num", "yoggies_url", "category", "name_cz", "name_nl",
    "description_cz", "description_nl", "ingredients_cz", "dosage_cz",
    "icons_text", "icons_images", "has_calculator",
    "image_main", "image_2", "image_3",
    "doglife_url", "doglife_image", "match_status", "match_score",
]

# Shopify product_type / tag -> the category the site builder expects.
CATEGORY_RULES = [
    ("yoghurt", ["jogurt", "yoghurt"]),
    ("barf", ["b.a.r.f", "barf", "syrove maso", "priloha k b.a.r.f"]),
    ("blikvoer", ["konzerva", "blik"]),
    ("gekookt", ["varena potrava", "vařená", "pastika", "paštika", "omacka", "omáčka"]),
    ("brokken", ["granule", "vzorek", "minigranule"]),
    ("snacks", ["pamlsek", "pamlsky"]),
    ("supplementen", ["doplnek potravy", "doplněk", "olej", "probiotik"]),
    ("speelgoed", ["hracka", "hračka", "miska", "lizaci podlozka", "vybaveni",
                   "hygiena", "obleceni", "kniha", "darkovy poukaz",
                   "treningova pomucka", "plnici hracka", "stelivo"]),
]


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def categorise(item: dict) -> str:
    haystack = norm(f"{item.get('product_type','')} {' '.join(item.get('tags') or [])} "
                    f"{item.get('name_cz','')}")
    for category, needles in CATEGORY_RULES:
        if any(norm(n) in haystack for n in needles):
            return category
    return "supplementen"


def first_nonempty(*values):
    """First value that is not None and not blank, else "" ."""
    for value in values:
        if str(value or "").strip():
            return value
    return ""


def read_existing(source: Path) -> tuple[dict, dict]:
    """Return (by_url, by_name_cz) of rows worth preserving from an old master.

    Keyed on the Czech name rather than the Dutch one: name_nl is exactly the
    field that may be missing or untranslated in the file being rebuilt, so it
    cannot serve as the fallback key for finding a previous translation.
    """
    if not source.exists():
        return {}, {}
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = {str(h): i for i, h in enumerate(next(rows)) if h}

    by_url, by_name = {}, {}
    for row in rows:
        def get(col):
            i = header.get(col)
            return row[i] if i is not None and i < len(row) else None

        keep = {
            "ingredients_cz": get("ingredients_cz"),
            "dosage_cz": get("dosage_cz"),
            "icons_text": get("icons_text"),
            "icons_images": get("icons_images"),
            "has_calculator": get("has_calculator"),
            # Translated content - costly to regenerate, so it survives a
            # rebuild that runs without a fresh translation pass.
            "name_nl": get("name_nl"),
            "description_nl": get("description_nl"),
        }
        url = str(get("yoggies_url") or "").rstrip("/")
        name_cz = norm(get("name_cz"))
        if url:
            by_url[url] = keep
        if name_cz:
            by_name[name_cz] = keep
    wb.close()
    return by_url, by_name


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--preserve-from", metavar="XLSX",
                    help="read preserved columns from this workbook instead of "
                         "the current master (e.g. a backup taken before a "
                         "rebuild that dropped translations)")
    args = ap.parse_args()

    preserve_src = Path(args.preserve_from).expanduser() if args.preserve_from else MASTER
    if args.preserve_from and not preserve_src.exists():
        sys.exit(f"--preserve-from: {preserve_src} not found")

    src = TRANSLATED if TRANSLATED.exists() else CATALOGUE
    if not src.exists():
        sys.exit("No catalogue found. Run pipe/1_fetch_yoggies.py first.")
    products = json.loads(src.read_text(encoding="utf-8"))

    matches = json.loads(MATCHES.read_text(encoding="utf-8")) if MATCHES.exists() else {}
    if not matches:
        print("Warning: no doglife_matches.json - purchase links will be empty.\n")

    by_url, by_name = read_existing(preserve_src)
    print(f"catalogue        : {len(products)} products (from {src.name})")
    print(f"preserved rows   : {len(by_url)} from {preserve_src.name}")
    print(f"doglife matches  : {len(matches)}\n")

    rows, stats = [], Counter()
    for n, item in enumerate(products, 1):
        url = item["url"].rstrip("/")
        # URL first, Czech name second: a renamed product keeps its old row.
        kept = by_url.get(url) or by_name.get(norm(item.get("name_cz"))) or {}
        if kept:
            stats["preserved"] += 1

        match = matches.get(item["handle"], {})
        status = match.get("status", "NOT_ON_DOGLIFE")
        stats[status] += 1

        images = item.get("images") or []
        category = categorise(item)
        stats[f"cat:{category}"] += 1

        dosage = kept.get("dosage_cz")
        has_calc = bool(kept.get("has_calculator")) and bool(dosage)

        # Fresh translation wins; otherwise reuse the previous one; otherwise
        # fall back to Czech so the field is never blank.
        name_nl = first_nonempty(
            item.get("name_nl"), kept.get("name_nl"), item.get("name_cz"))
        description_nl = first_nonempty(
            item.get("description_nl"), kept.get("description_nl"),
            item.get("description_cz"))

        for field, fresh, old in (
            ("name", item.get("name_nl"), kept.get("name_nl")),
            ("desc", item.get("description_nl"), kept.get("description_nl")),
        ):
            if str(fresh or "").strip():
                stats[f"{field}:translated"] += 1
            elif str(old or "").strip():
                stats[f"{field}:preserved"] += 1
            else:
                stats[f"{field}:czech"] += 1

        rows.append([
            n,
            item["url"],
            category,
            item.get("name_cz", ""),
            name_nl,
            item.get("description_cz", ""),
            description_nl,
            kept.get("ingredients_cz"),
            dosage,
            kept.get("icons_text"),
            kept.get("icons_images"),
            has_calc,
            item.get("image_main", ""),
            images[1] if len(images) > 1 else None,
            images[2] if len(images) > 2 else None,
            match.get("doglife_url", ""),
            match.get("doglife_image", ""),
            status,
            1.0 if status.startswith("MATCH") else 0.0,
        ])

    print("Categories")
    for k, v in sorted(stats.items()):
        if k.startswith("cat:"):
            print(f"  {v:4d}  {k[4:]}")
    print("\nDogLife status")
    for k in ("MATCH", "MATCH_CONFIRMED", "REVIEW", "NOT_ON_DOGLIFE"):
        if stats[k]:
            print(f"  {stats[k]:4d}  {k}")
    print(f"\npreserved ingredient/calculator data on {stats['preserved']} products")
    print("\nDutch text sources")
    for field, label in (("name", "name_nl"), ("desc", "description_nl")):
        print(f"  {label:16s} {stats[f'{field}:translated']:4d} fresh translation"
              f" | {stats[f'{field}:preserved']:4d} preserved"
              f" | {stats[f'{field}:czech']:4d} Czech fallback")

    if args.dry_run:
        print("\nDry run - nothing written.")
        return 0

    if MASTER.exists():
        backup = MASTER.with_name(
            f"READY_FOR_BUILD_FINAL.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
        shutil.copy2(MASTER, backup)
        print(f"\nbacked up old master -> {backup.name}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    wb.save(MASTER)
    print(f"Wrote {MASTER.name} with {len(rows)} products")
    print("\nNext: python pipe/audit.py")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
