"""
Diagnose variant detection: shows what's in the Excel and what would/wouldn't merge.

Writes the same report to diagnose_output.txt (UTF-8).

Do NOT capture this with a shell redirect (`python diagnose_variants.py > out.txt`)
— PowerShell writes UTF-16 and mangles the Czech characters. The script writes
its own report file so the console codepage never matters.
"""
import atexit
import re
import sys
from pathlib import Path
from collections import defaultdict
import openpyxl

ROOT = Path(__file__).resolve().parent
OUT_FILE = ROOT / "diagnose_output.txt"


class _Tee:
    """Fan stdout out to the console and to the report file simultaneously."""

    def __init__(self, *streams):
        self._streams = streams

    def write(self, data):
        for s in self._streams:
            if not s.closed:
                s.write(data)
        return len(data)

    def flush(self):
        # The interpreter flushes sys.stdout during shutdown, after atexit has
        # already closed the report file — skip whatever is gone by then.
        for s in self._streams:
            if not s.closed:
                s.flush()


# Czech characters must survive on a cp850/cp1252 console too.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, ValueError):  # pragma: no cover - old/oddly wrapped stdout
    pass

_report = OUT_FILE.open("w", encoding="utf-8", newline="\n")
atexit.register(_report.close)
sys.stdout = _Tee(sys.stdout, _report)

SIZE_PATTERNS = [
    re.compile(r'^(\d+(?:[.,]\d+)?)\s*(kg|g|ml|l)\s+', re.IGNORECASE),
    re.compile(r'\s+(\d+(?:[.,]\d+)?)\s*(kg|g|ml|l)\s*$', re.IGNORECASE),
]

def extract_size(name):
    for pat in SIZE_PATTERNS:
        m = pat.search(name)
        if m:
            qty = m.group(1).strip()
            unit = m.group(2).lower()
            stripped = (name[:m.start()] + " " + name[m.end():]).strip()
            stripped = re.sub(r'\s+', ' ', stripped)
            return stripped, f"{qty} {unit}"
    return name, None

xlsx_path = ROOT / "READY_FOR_BUILD_FINAL.xlsx"
if not xlsx_path.exists():
    raise SystemExit(f"X {xlsx_path} not found.")

wb = openpyxl.load_workbook(xlsx_path, data_only=True)
ws = wb.active
hdr = [str(c.value or "").strip() for c in ws[1]]

def col(name):
    return hdr.index(name) if name in hdr else -1

name_i = col("name_nl")
cat_i  = col("category")
url_i  = col("doglife_url")

if name_i < 0 or cat_i < 0:
    print(f"Available columns: {hdr}")
    raise SystemExit("X Missing 'name_nl' or 'category' column")

products_by_cat = defaultdict(list)
for row in ws.iter_rows(min_row=2, values_only=True):
    name = str(row[name_i] or "").strip()
    cat  = str(row[cat_i] or "").strip()
    url  = str(row[url_i] or "").strip() if url_i >= 0 else ""
    if name:
        products_by_cat[cat].append((name, url))

print(f"\nTotal products in Excel: {sum(len(v) for v in products_by_cat.values())}")
print(f"Categories: {len(products_by_cat)}")
print("\n" + "=" * 80)
print("VARIANT DETECTION BY CATEGORY")
print("=" * 80)

for cat, products in sorted(products_by_cat.items()):
    print(f"\n>> {cat} ({len(products)} products)")
    print("-" * 80)

    by_base = defaultdict(list)
    no_size = []
    for name, url in products:
        base, size = extract_size(name)
        if size:
            by_base[base.lower()].append((name, size, url, base))
        else:
            no_size.append((name, url))

    groups = {k: v for k, v in by_base.items() if len(v) >= 2}
    if groups:
        print(f"  [+] {len(groups)} variant group(s):")
        for items in groups.values():
            base_name = items[0][3]
            print(f"    * {base_name}")
            for n, s, u, _ in items:
                tail = u.split("/")[-1][:50] if u else "(no DogLife URL)"
                print(f"        [{s:>7}] {n}")
                print(f"                  -> {tail}")

    singletons_with_size = [v[0] for v in by_base.values() if len(v) == 1]
    if singletons_with_size:
        print(f"  [.] {len(singletons_with_size)} product(s) with size but NO sibling:")
        for n, s, u, _ in singletons_with_size[:10]:
            print(f"      [{s}] {n}")
        if len(singletons_with_size) > 10:
            print(f"      ... +{len(singletons_with_size)-10} more")

    if no_size:
        print(f"  [.] {len(no_size)} product(s) without any size in name:")
        for n, u in no_size[:10]:
            print(f"      {n}")
        if len(no_size) > 10:
            print(f"      ... +{len(no_size)-10} more")

print("\n" + "=" * 80)
print("DONE")
print("=" * 80)
